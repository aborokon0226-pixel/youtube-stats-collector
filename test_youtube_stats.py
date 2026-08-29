import unittest
from datetime import date

from openpyxl import Workbook

from youtube_stats import (
    KEY_CONTENT_VIEWS_ROW,
    PULLING_CONTENT_VIEWS_ROW,
    compute_daily_average_views,
    compute_subscriber_conversion_rate,
    dedupe_ids,
    ensure_week_column,
    estimate_48_hour_views,
    extract_video_id,
    get_all_playlist_video_ids,
    get_previous_subscriber_snapshot,
    links_to_video_ids,
    parse_channel_reference,
    record_subscriber_snapshot,
    week_column_index,
    week_date_range,
    week_of_month,
    write_view_counts,
)


class TestExtractVideoId(unittest.TestCase):
    def test_standard_watch_url(self):
        url = "https://www.youtube.com/watch?v=dQw4w9WgXcQ"
        self.assertEqual(extract_video_id(url), "dQw4w9WgXcQ")

    def test_short_url(self):
        url = "https://youtu.be/dQw4w9WgXcQ"
        self.assertEqual(extract_video_id(url), "dQw4w9WgXcQ")

    def test_shorts_url(self):
        url = "https://www.youtube.com/shorts/dQw4w9WgXcQ"
        self.assertEqual(extract_video_id(url), "dQw4w9WgXcQ")

    def test_url_with_extra_params(self):
        url = "https://www.youtube.com/watch?v=dQw4w9WgXcQ&t=30s"
        self.assertEqual(extract_video_id(url), "dQw4w9WgXcQ")

    def test_invalid_url_returns_none(self):
        url = "https://www.example.com/not-a-youtube-link"
        self.assertIsNone(extract_video_id(url))


class TestDedupeIds(unittest.TestCase):
    def test_removes_duplicates_preserving_order(self):
        ids = ["a", "b", "a", "c", "b"]
        self.assertEqual(dedupe_ids(ids), ["a", "b", "c"])

    def test_no_duplicates_returns_same_list(self):
        ids = ["a", "b", "c"]
        self.assertEqual(dedupe_ids(ids), ["a", "b", "c"])

    def test_empty_list(self):
        self.assertEqual(dedupe_ids([]), [])


class TestLinksToVideoIds(unittest.TestCase):
    def test_splits_valid_and_invalid_links(self):
        links = [
            "https://youtu.be/dQw4w9WgXcQ",
            "not-a-link",
            "https://youtu.be/dQw4w9WgXcQ",  # 중복
        ]
        video_ids, invalid_links = links_to_video_ids(links)
        self.assertEqual(video_ids, ["dQw4w9WgXcQ"])
        self.assertEqual(invalid_links, ["not-a-link"])


class TestWeekOfMonth(unittest.TestCase):
    def test_first_week(self):
        self.assertEqual(week_of_month(date(2026, 8, 1)), 1)
        self.assertEqual(week_of_month(date(2026, 8, 7)), 1)

    def test_third_week(self):
        self.assertEqual(week_of_month(date(2026, 8, 15)), 3)
        self.assertEqual(week_of_month(date(2026, 8, 19)), 3)
        self.assertEqual(week_of_month(date(2026, 8, 21)), 3)

    def test_last_partial_week(self):
        self.assertEqual(week_of_month(date(2026, 8, 31)), 5)


class TestWeekDateRange(unittest.TestCase):
    def test_third_week_of_august(self):
        start, end = week_date_range(date(2026, 8, 19))
        self.assertEqual(start, date(2026, 8, 15))
        self.assertEqual(end, date(2026, 8, 21))

    def test_last_week_clamps_to_month_end(self):
        start, end = week_date_range(date(2026, 8, 31))
        self.assertEqual(start, date(2026, 8, 29))
        self.assertEqual(end, date(2026, 8, 31))


class TestParseChannelReference(unittest.TestCase):
    def test_channel_url(self):
        url = "https://www.youtube.com/channel/UCxxxxxxxxxxxxxxxxxxxxxx"
        self.assertEqual(parse_channel_reference(url), {"id": "UCxxxxxxxxxxxxxxxxxxxxxx"})

    def test_handle_url(self):
        url = "https://www.youtube.com/@example_channel"
        self.assertEqual(parse_channel_reference(url), {"forHandle": "@example_channel"})

    def test_bare_handle(self):
        self.assertEqual(parse_channel_reference("@example_channel"), {"forHandle": "@example_channel"})

    def test_bare_channel_id(self):
        channel_id = "UC" + "x" * 22
        self.assertEqual(parse_channel_reference(channel_id), {"id": channel_id})

    def test_plain_name_treated_as_handle(self):
        self.assertEqual(parse_channel_reference("example_channel"), {"forHandle": "@example_channel"})


class _FakePlaylistItemsRequest:
    def __init__(self, pages, page_token):
        self._pages = pages
        self._page_token = page_token

    def execute(self):
        return self._pages[self._page_token or 0]


class _FakeYoutubeClient:
    """playlistItems().list().execute() 체이닝만 흉내내는 테스트용 스텁."""

    def __init__(self, pages):
        self._pages = pages  # {page_index: response_dict}

    def playlistItems(self):
        return self

    def list(self, part, playlistId, maxResults, pageToken):
        index = 0 if pageToken is None else pageToken
        return _FakePlaylistItemsRequest(self._pages, index)


class TestGetAllPlaylistVideoIds(unittest.TestCase):
    def test_paginates_through_all_pages(self):
        pages = {
            0: {
                "items": [{"contentDetails": {"videoId": "a"}}, {"contentDetails": {"videoId": "b"}}],
                "nextPageToken": 1,
            },
            1: {
                "items": [{"contentDetails": {"videoId": "c"}}],
            },
        }
        youtube = _FakeYoutubeClient(pages)
        self.assertEqual(get_all_playlist_video_ids(youtube, "PLxxxx"), ["a", "b", "c"])


class TestWeekColumnIndex(unittest.TestCase):
    def test_week_one_is_column_c(self):
        self.assertEqual(week_column_index(1), 3)

    def test_week_three_is_column_e(self):
        self.assertEqual(week_column_index(3), 5)


class TestEnsureWeekColumn(unittest.TestCase):
    def test_creates_header_for_new_week(self):
        sheet = Workbook().active
        column = ensure_week_column(sheet, 1)
        self.assertEqual(column, 3)
        self.assertEqual(sheet.cell(row=1, column=3).value, "1주차 현황")

    def test_keeps_existing_header(self):
        sheet = Workbook().active
        sheet.cell(row=1, column=3, value="1주차 현황(수정됨)")
        ensure_week_column(sheet, 1)
        self.assertEqual(sheet.cell(row=1, column=3).value, "1주차 현황(수정됨)")


class TestWriteViewCounts(unittest.TestCase):
    def test_writes_to_correct_rows(self):
        sheet = Workbook().active
        write_view_counts(sheet, week_col=3, key_content_views=100, pulling_content_views=200)
        self.assertEqual(sheet.cell(row=KEY_CONTENT_VIEWS_ROW, column=3).value, 100)
        self.assertEqual(sheet.cell(row=PULLING_CONTENT_VIEWS_ROW, column=3).value, 200)


class TestDailyAndFortyEightHourViews(unittest.TestCase):
    def test_daily_average(self):
        self.assertAlmostEqual(compute_daily_average_views(700), 100.0)

    def test_forty_eight_hour_estimate(self):
        self.assertAlmostEqual(estimate_48_hour_views(100.0), 200.0)


class TestComputeSubscriberConversionRate(unittest.TestCase):
    def test_normal_case(self):
        self.assertAlmostEqual(compute_subscriber_conversion_rate(20, 1000), 2.0)

    def test_zero_views_returns_zero(self):
        self.assertEqual(compute_subscriber_conversion_rate(5, 0), 0.0)

    def test_negative_growth(self):
        self.assertAlmostEqual(compute_subscriber_conversion_rate(-8, 500), -1.6)


class TestSubscriberSnapshot(unittest.TestCase):
    def test_no_snapshot_sheet_returns_none(self):
        workbook = Workbook()
        self.assertIsNone(get_previous_subscriber_snapshot(workbook, week=3, channel_id="UCabc"))

    def test_records_and_finds_previous_week(self):
        workbook = Workbook()
        record_subscriber_snapshot(workbook, week=1, subscriber_count=100, channel_id="UCabc")
        record_subscriber_snapshot(workbook, week=2, subscriber_count=120, channel_id="UCabc")
        self.assertEqual(get_previous_subscriber_snapshot(workbook, week=3, channel_id="UCabc"), 120)

    def test_ignores_future_weeks(self):
        workbook = Workbook()
        record_subscriber_snapshot(workbook, week=1, subscriber_count=100, channel_id="UCabc")
        record_subscriber_snapshot(workbook, week=5, subscriber_count=200, channel_id="UCabc")
        self.assertEqual(get_previous_subscriber_snapshot(workbook, week=3, channel_id="UCabc"), 100)

    def test_rerecording_same_week_updates_value(self):
        workbook = Workbook()
        record_subscriber_snapshot(workbook, week=1, subscriber_count=100, channel_id="UCabc")
        record_subscriber_snapshot(workbook, week=1, subscriber_count=150, channel_id="UCabc")
        self.assertEqual(get_previous_subscriber_snapshot(workbook, week=2, channel_id="UCabc"), 150)

    def test_ignores_snapshots_from_a_different_channel(self):
        workbook = Workbook()
        record_subscriber_snapshot(workbook, week=1, subscriber_count=100, channel_id="UCold")
        self.assertIsNone(get_previous_subscriber_snapshot(workbook, week=3, channel_id="UCnew"))


if __name__ == "__main__":
    unittest.main()
